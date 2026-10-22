"""Mensajes masivos por WhatsApp (Blueprint): campañas, historial de envíos y respuestas.

El envío real lo hace el servicio Node de la carpeta whatsapp-service/ (whatsapp-web.js), que
corre aparte; esta app le entrega los mensajes y él avisa de cada resultado/respuesta por webhook.
"""
import secrets
from datetime import date, datetime, timedelta

from flask import Blueprint, Response, abort, flash, g, jsonify, redirect, render_template, request, url_for

import config
import whatsapp_db as W
from importar_inventario import _alias
from importar_personas import _matriz, filas_de
from nomina_excel import _hoja
from seguridad import requiere_permiso, tiene

bp = Blueprint("whatsapp", __name__, url_prefix="/whatsapp")
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WEBHOOKS = ("whatsapp.webhook_estado", "whatsapp.webhook_respuesta", "whatsapp.webhook_archivo", "whatsapp.webhook_latido", "whatsapp.webhook_lid")

COLUMNAS = [("fecha", "Fecha", "2026-10-15"), ("manifiesto", "Manifiesto", "MF-123456"),
            ("placa", "Placa", "ABC123"), ("celular", "Celular", "3001234567"), ("contacto", "Contacto", "Juan Pérez")]
ALIAS = _alias(
    ("fecha", ["fecha", "fechamanifiesto", "fechadespacho", "fechaviaje"]),
    ("manifiesto", ["manifiesto", "mft", "nummanifiesto", "numeromanifiesto"]),
    ("placa", ["placa", "placavehiculo", "vehiculo"]),
    ("celular", ["celular", "telefono", "whatsapp", "numero", "movil", "tel"]),
    ("contacto", ["contacto", "nombre", "conductor", "nombrecontacto", "nombreconductor"]))
MAX_FILAS = 1000


def _cfg(nombre, defecto=None):
    return getattr(config, nombre, defecto)


# ---------------------------------------------------------------- acceso
@bp.before_request
def acceso():
    if request.endpoint in WEBHOOKS:
        enviado = request.headers.get("X-WA-Token", "").encode()
        esperado = (_cfg("WA_TOKEN") or "").encode()
        if not esperado or not secrets.compare_digest(enviado, esperado):
            abort(403)
        return
    if not g.get("usuario"):
        return redirect(url_for("login", siguiente=request.path))
    if not tiene("whatsapp.consultar"):
        abort(403)


@bp.app_context_processor
def _ctx():
    return dict(WA_ESTADOS=W.ESTADOS)


# ---------------------------------------------------------------- servicio Node
# El servicio (whatsapp-service/) consulta a esta app cada pocos segundos (POST /whatsapp/webhook/latido):
# informa su estado y recibe los mensajes "en cola". No hace falta abrir ningún puerto hacia el servicio.
def estado_servicio():
    return W.estado_servicio()


# ---------------------------------------------------------------- pantallas
@bp.route("/")
def lista():
    est, error = estado_servicio()
    return render_template("whatsapp/lista.html", campanas=W.campanas(), est=est, error_servicio=error)


@bp.route("/conexion")
@requiere_permiso("whatsapp.enviar")
def conexion():
    est, error = estado_servicio()
    return render_template("whatsapp/conexion.html", est=est, error_servicio=error)


@bp.route("/conexion/reiniciar", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def reiniciar_servicio():
    """Pide al servicio de WhatsApp que se reinicie. Solo llega si el servicio sigue enviando señal."""
    est, error = estado_servicio()
    if error:
        flash("El servicio no está enviando señal, así que no puede reiniciarse desde aquí. Enciéndelo en el PC de la "
              "oficina con «iniciar_whatsapp.bat» (carpeta whatsapp-service).")
    else:
        W.pedir_reinicio()
        flash("Reinicio pedido. El servicio se apaga y vuelve a conectarse solo en 1-2 minutos; recarga la página para ver el estado.")
    return redirect(url_for("whatsapp.conexion"))


@bp.route("/conexion/cerrar", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def cerrar_sesion():
    W.pedir_cierre()
    flash("Se pidió cerrar la sesión de WhatsApp. En unos segundos aparecerá un QR nuevo para volver a conectar.")
    return redirect(url_for("whatsapp.conexion"))


def _fecha(texto):
    """(date|None, valida). Acepta aaaa-mm-dd, dd/mm/aaaa, dd-mm-aaaa y el número de serie de Excel."""
    t = (texto or "").strip().split(" ")[0]
    if not t:
        return None, True
    if t.isdigit() and 20000 < int(t) < 80000:
        return date(1899, 12, 30) + timedelta(days=int(t)), True
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t, fmt).date(), True
        except ValueError:
            pass
    return None, False


def leer_archivo(archivo):
    """Agrupa por celular: devuelve ([{telefono, contacto, items:[{fecha, manifiesto, placa}]}], errores)."""
    hojas, error = _matriz(archivo.filename, archivo.read())
    if error:
        return None, [error]
    filas, error = filas_de(next(iter(hojas.values())), ALIAS, "celular", [c for c, _, _ in COLUMNAS])
    if error:
        return None, [error]
    if len(filas) > MAX_FILAS:
        return None, ["El archivo tiene más de %d filas. Divídelo en partes." % MAX_FILAS]
    pais = _cfg("WA_PAIS", "57")
    grupos, errores, vistos = {}, [], set()
    for f in filas:
        tel, e = W.normalizar_telefono(f["celular"], pais)
        if e:
            errores.append("Fila %d: %s" % (f["_fila"], e.replace("el teléfono", "el celular").replace("Teléfono", "Celular")))
            continue
        man = f["manifiesto"].strip()[:60]
        if not man:
            errores.append("Fila %d: falta el manifiesto" % f["_fila"])
            continue
        fecha, valida = _fecha(f["fecha"])
        if not valida:
            errores.append("Fila %d: la fecha «%s» no es válida (usa aaaa-mm-dd o dd/mm/aaaa)" % (f["_fila"], f["fecha"]))
            continue
        if (tel, man.lower()) in vistos:
            errores.append("Fila %d: el manifiesto %s está repetido para ese celular" % (f["_fila"], man))
            continue
        vistos.add((tel, man.lower()))
        g = grupos.setdefault(tel, {"telefono": tel, "contacto": None, "items": []})
        g["contacto"] = g["contacto"] or f["contacto"].strip()[:150] or None
        g["items"].append({"fecha": fecha, "manifiesto": man, "placa": f["placa"].strip()[:30] or None})
    return list(grupos.values()), errores


@bp.route("/nueva", methods=["GET", "POST"])
@requiere_permiso("whatsapp.enviar")
def nueva():
    v = {"nombre": "Cumplidos pendientes " + date.today().strftime("%d/%m/%Y"), "plantilla": W.PLANTILLA_DEFECTO,
         "segundo": False, "plantilla2": W.PLANTILLA_SEGUNDO}
    errores = []
    if request.method == "POST":
        v = {"nombre": (request.form.get("nombre") or "").strip()[:150],
             "plantilla": (request.form.get("plantilla") or "").strip()[:2000],
             "segundo": request.form.get("segundo") == "1",
             "plantilla2": (request.form.get("plantilla2") or "").strip()[:2000]}
        arch = request.files.get("archivo")
        if not v["nombre"]:
            errores.append("Escribe un nombre para la campaña.")
        if not v["plantilla"]:
            errores.append("Escribe el texto del mensaje.")
        if v["segundo"] and not v["plantilla2"]:
            errores.append("Escribe el texto del segundo aviso o desmarca «Habilitar segundo aviso».")
        if not arch or not arch.filename:
            errores.append("Selecciona el archivo .xlsx o .csv con los pendientes.")
        if not errores:
            contactos, errores = leer_archivo(arch)
            if not errores and not contactos:
                errores = ["El archivo no tiene filas para enviar."]
            if not errores:
                cid = W.crear_campana(v["nombre"], v["plantilla"], v["plantilla2"] if v["segundo"] else None,
                                      g.usuario["nombre"], contactos)
                flash("Campaña creada: %d contactos con %d manifiestos. Revisa los mensajes y pulsa «Enviar» cuando estés listo."
                      % (len(contactos), sum(len(c["items"]) for c in contactos)))
                return redirect(url_for("whatsapp.detalle", cid=cid))
    return render_template("whatsapp/nueva.html", v=v, errores=errores[:60], mas_errores=max(0, len(errores) - 60),
                           columnas=COLUMNAS, defecto2=W.PLANTILLA_SEGUNDO)


@bp.route("/plantilla")
@requiere_permiso("whatsapp.enviar")
def plantilla():
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    wb.remove(wb.active)
    ws = _hoja(wb, "Pendientes", [t for _, t, _ in COLUMNAS], [
        [date(2026, 10, 15), "MF-123456", "ABC123", "3001234567", "Juan Pérez"],
        [date(2026, 10, 17), "MF-123789", "XYZ987", "3001234567", "Juan Pérez"],
        [date(2026, 10, 16), "MF-124000", "DEF456", "3109876543", "Ana Gómez"]])
    for fila in range(2, ws.max_row + 1):
        ws.cell(row=fila, column=1).number_format = "dd/mm/yyyy"
        for col in (2, 3, 4):
            ws.cell(row=fila, column=col).number_format = "@"
    hoja = wb.create_sheet("Instrucciones")
    for linea in ("Una fila por manifiesto. La fila 1 son los títulos: no los cambies. Borra las filas de ejemplo.",
                  "Un mismo contacto (celular) puede tener varios manifiestos y placas: repite el celular en cada fila y "
                  "recibirá UN solo mensaje con todos.",
                  "Obligatorios: Manifiesto y Celular. Fecha: aaaa-mm-dd o dd/mm/aaaa. Placa y Contacto se muestran en el mensaje.",
                  "Celular: 10 dígitos de celular colombiano (3001234567) o con indicativo (573001234567).",
                  "Formatea Manifiesto, Placa y Celular como TEXTO para no perder ceros a la izquierda."):
        hoja.append([linea])
    hoja.column_dimensions["A"].width = 110
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype=EXCEL_MIME,
                    headers={"Content-Disposition": 'attachment; filename="plantilla_whatsapp_pendientes.xlsx"'})


def campana_o_404(cid):
    return W.campana(cid) or abort(404)


@bp.route("/<int:cid>")
def detalle(cid):
    c = campana_o_404(cid)
    estado = request.args.get("estado")
    if estado != "finalizado" and estado not in W.ESTADOS:
        estado = None
    est, error = estado_servicio()
    msgs = W.mensajes(cid, estado)
    conteo, finalizados = {}, 0
    for m in W.mensajes(cid):
        if m["finalizado"]:                     # lo finalizado no cuenta como pendiente/en cola/error
            finalizados += m["aviso"] == 1
        else:
            conteo[m["estado"]] = conteo.get(m["estado"], 0) + 1
    sin, con = W.candidatos_segundo(cid) if c["plantilla2"] else ([], [])
    return render_template("whatsapp/campana.html", c=c, msgs=msgs, conteo=conteo, filtro=estado, est=est,
                           error_servicio=error, resp=W.respuestas_de_campana(cid), items=W.items_por_telefono(cid),
                           segundo_sin=len(sin), segundo_con=len(con), finalizados=finalizados)


@bp.route("/<int:cid>/enviar", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def enviar(cid):
    campana_o_404(cid)
    est, error = estado_servicio()
    if error or est.get("estado") != "listo":
        flash(error or "WhatsApp no está conectado. Ve a «Conexión» y escanea el QR.")
        return redirect(url_for("whatsapp.detalle", cid=cid))
    estados = ("pendiente", "error") if request.form.get("reintentar") == "1" else ("pendiente",)
    pendientes = [m for m in W.mensajes(cid) if m["estado"] in estados and not m["finalizado"]]
    if not pendientes:
        flash("No hay mensajes por enviar.")
        return redirect(url_for("whatsapp.detalle", cid=cid))
    ids = [m["id"] for m in pendientes]
    W.marcar_en_cola(ids)
    flash("%d mensajes en cola. El servicio los toma en unos segundos y los envía uno a uno con pausas aleatorias; "
          "esta página se puede cerrar." % len(ids))
    return redirect(url_for("whatsapp.detalle", cid=cid))


@bp.route("/<int:cid>/segundo-aviso", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def segundo_aviso(cid):
    c = campana_o_404(cid)
    if not c["plantilla2"]:
        flash("Esta campaña no tiene el segundo aviso habilitado.")
        return redirect(url_for("whatsapp.detalle", cid=cid))
    est, error = estado_servicio()
    if error or est.get("estado") != "listo":
        flash(error or "WhatsApp no está conectado. Ve a «Conexión» y escanea el QR.")
        return redirect(url_for("whatsapp.detalle", cid=cid))
    n = W.crear_segundo_aviso(cid, c["plantilla2"], request.form.get("incluir_respondieron") == "1")
    flash("%d segundos avisos en cola." % n if n else "No hay contactos a quienes enviar el segundo aviso.")
    return redirect(url_for("whatsapp.detalle", cid=cid))


@bp.route("/<int:cid>/finalizar", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def finalizar(cid):
    """Finaliza (o reabre) uno o varios contactos de la campaña."""
    campana_o_404(cid)
    ids = [int(x) for x in request.form.getlist("ids") if x.isdigit()]
    volver = request.form.get("estado")
    volver = volver if volver == "finalizado" or volver in W.ESTADOS else None
    if not ids:
        flash("Selecciona al menos un contacto.")
    else:
        reabrir = request.form.get("accion") == "reabrir"
        n = W.finalizar(cid, ids, g.usuario["nombre"], not reabrir)
        flash(("%d contacto(s) reabierto(s)." if reabrir else
               "%d contacto(s) finalizado(s): no recibirán más mensajes ni el segundo aviso.") % n)
    return redirect(url_for("whatsapp.detalle", cid=cid, estado=volver))


@bp.route("/<int:cid>/reponer", methods=["POST"])
@requiere_permiso("whatsapp.enviar")
def reponer(cid):
    """Cancela lo que sigue en cola (los que ya salieron no se tocan)."""
    campana_o_404(cid)
    W.reponer_en_cola(cid)
    flash("Los mensajes en cola volvieron a «Pendiente».")
    return redirect(url_for("whatsapp.detalle", cid=cid))


@bp.route("/<int:cid>/exportar")
def exportar(cid):
    from io import BytesIO
    from openpyxl import Workbook
    c = campana_o_404(cid)
    fmt = lambda d: d.strftime("%Y-%m-%d %H:%M") if d else ""
    msgs = W.mensajes(cid)
    wb = Workbook()
    wb.remove(wb.active)
    _hoja(wb, "Envíos", ["Contacto", "Celular", "Aviso", "Manifiestos", "Estado", "Enviado", "Entregado", "Leído",
                         "Respuestas", "Última respuesta", "Error", "Finalizado"],
          [[m["contacto"] or "", m["telefono"], m["aviso"], m["manifiestos"] or m["manifiesto"] or "",
            W.ESTADOS.get(m["estado"], m["estado"]), fmt(m["enviado_en"]), fmt(m["entregado_en"]), fmt(m["leido_en"]),
            m["n_resp"], m["ultima_resp"] or "", m["error"] or "",
            ("Sí · %s · %s" % (m["finalizado_por"] or "", fmt(m["finalizado_en"]))) if m["finalizado"] else ""]
           for m in msgs])
    _hoja(wb, "Manifiestos", ["Contacto", "Celular", "Fecha", "Manifiesto", "Placa"],
          [[(i["contacto"] if "contacto" in i else "") or "", tel, i["fecha"].strftime("%Y-%m-%d") if i["fecha"] else "",
            i["manifiesto"], i["placa"] or ""] for tel, lista in W.items_por_telefono(cid).items() for i in lista])
    _hoja(wb, "Respuestas", ["Recibida", "Contacto", "Celular", "Manifiestos", "Respuesta", "Archivo en OneDrive", "Enlace"],
          [[fmt(r["recibida_en"]), r["contacto"] or "", r["telefono"], r["manifiesto"] or "",
            r["texto"] or "(%s)" % r["tipo"], r["archivo_nombre"] or "", r["archivo_url"] or ""]
           for r in W.respuestas_de_campana(cid)])
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype=EXCEL_MIME, headers={
        "Content-Disposition": 'attachment; filename="CC_Gestion_whatsapp_%d_%s.xlsx"' % (c["id"], date.today().strftime("%Y%m%d"))})


@bp.app_template_filter("tel")
def _tel(n):
    """573001234567 -> +57 300 123 4567"""
    n = str(n or "")
    if len(n) == 12 and n.startswith("57"):
        return "+57 %s %s %s" % (n[2:5], n[5:8], n[8:])
    return "+" + n if n else ""


@bp.route("/respuestas")
def todas_respuestas():
    f = {k: (request.args.get(k) or "").strip()[:80] for k in ("placa", "contacto", "manifiesto", "celular")}
    f["campana"] = int(request.args["campana"]) if (request.args.get("campana") or "").isdigit() else None
    filas = W.respuestas(f)
    return render_template("whatsapp/respuestas.html", filas=filas, f=f, campanas=W.campanas(),
                           filtrado=any(f.values()))


# ---------------------------------------------------------------- webhooks (los llama el servicio Node)
@bp.route("/webhook/estado", methods=["POST"])
def webhook_estado():
    d = request.get_json(silent=True) or {}
    if "ack" in d:
        W.confirmacion(d.get("wa_id"), int(d["ack"]))
    elif d.get("id"):
        W.resultado_envio(int(d["id"]), bool(d.get("ok")), d.get("wa_id"), d.get("error"), str(d.get("chat") or "") or None)
    return jsonify(ok=True)


@bp.route("/webhook/respuesta", methods=["POST"])
def webhook_respuesta():
    d = request.get_json(silent=True) or {}
    tel = "".join(ch for ch in str(d.get("telefono", "")) if ch.isdigit())[:20]
    if not tel:
        return jsonify(ok=False), 400
    try:
        cuando = datetime.fromtimestamp(int(d.get("ts") or 0)) if d.get("ts") else W.ahora()
    except (ValueError, OSError):
        cuando = W.ahora()
    rid, m = W.guardar_respuesta(tel, d.get("texto"), str(d.get("tipo") or "chat"), d.get("wa_id"), d.get("citado_id"),
                                 cuando, con_archivo=bool(d.get("tiene_archivo")),
                                 chat_id=str(d.get("chat_id") or "") or None)
    # El servicio usa el manifiesto/cliente para nombrar y ubicar el adjunto en OneDrive
    return jsonify(ok=True, respuesta_id=rid, mensaje=({"manifiesto": m["manifiesto"], "cliente": m["cliente"]} if m else None))


@bp.route("/webhook/archivo", methods=["POST"])
def webhook_archivo():
    d = request.get_json(silent=True) or {}
    if not str(d.get("respuesta_id", "")).isdigit():
        return jsonify(ok=False), 400
    W.resultado_archivo(int(d["respuesta_id"]), d.get("estado"), d.get("nombre"), d.get("url"), d.get("error"))
    return jsonify(ok=True)


@bp.route("/webhook/latido", methods=["POST"])
def webhook_latido():
    """Señal periódica del servicio: estado/QR; responde con los mensajes por enviar y órdenes pendientes."""
    d = request.get_json(silent=True) or {}
    qr = d.get("qr") if isinstance(d.get("qr"), str) and d["qr"].startswith("data:image/png;base64,") else None
    comando, mensajes, vigentes, lids = W.latido(str(d.get("estado") or "desconocido")[:20], qr,
                                          str(d.get("numero") or "")[:30] or None, int(d.get("en_cola") or 0))
    return jsonify(ok=True, comando=comando, mensajes=mensajes, vigentes=vigentes, lids=lids,
                   ahora=W.ahora().strftime("%Y-%m-%d %H:%M:%S"))


@bp.route("/webhook/lid", methods=["POST"])
def webhook_lid():
    """El servicio tradujo un ID interno de WhatsApp a su celular real: se corrigen y ligan las respuestas."""
    d = request.get_json(silent=True) or {}
    lid = "".join(ch for ch in str(d.get("lid", "")) if ch.isdigit())[:20]
    tel = "".join(ch for ch in str(d.get("telefono", "")) if ch.isdigit())[:20]
    if not lid or not tel or lid == tel:
        return jsonify(ok=False), 400
    return jsonify(ok=True, corregidas=W.corregir_lid(lid, tel))
