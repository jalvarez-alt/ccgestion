"""Excel de nómina: resumen, conceptos, aportes del empleador y provisiones."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre",
         "Octubre", "Noviembre", "Diciembre"]


def _hoja(wb, nombre, cabecera, filas, dinero=()):
    ws = wb.create_sheet(nombre)
    ws.append(cabecera)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="254AA5")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    anchos = [len(h) + 2 for h in cabecera]
    for f in filas:
        ws.append(f)
        n = ws.max_row
        for i, v in enumerate(f, start=1):
            celda = ws.cell(row=n, column=i)
            if isinstance(v, str):
                if v.startswith(("=", "+", "-", "@")):
                    celda.data_type = "s"
                anchos[i - 1] = max(anchos[i - 1], min(len(v) + 2, 40))
            elif i - 1 in dinero:
                celda.number_format = "#,##0"
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(a, 10), 40)
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def generar_excel_nomina(liquidaciones):
    """liquidaciones: lista de dicts de nomina_db.consultar()."""
    wb = Workbook()
    wb.remove(wb.active)
    base = lambda l: [l["empleador"], l["anio"], MESES[l["mes"]], l["quincena"]]
    cab_base = ["Empleador", "Año", "Mes", "Quincena"]

    resumen = []
    for l in liquidaciones:
        c = l["calculo"]
        resumen.append(base(l) + [l["num_id"], l["nombre"], l.get("cargo") or "", l["dias_trabajados"],
                                  c["salario_devengado"], l["total_devengado"], l["total_deducciones"], l["neto"],
                                  l["ibc"], c.get("recobrable", 0)])
    _hoja(wb, "Resumen", cab_base + ["Documento", "Empleado", "Cargo", "Días trabajados", "Salario devengado",
                                      "Total devengado", "Total deducciones", "Neto a pagar", "IBC",
                                      "Recobrable EPS/ARL"], resumen, dinero=(8, 9, 10, 11, 12, 13))

    conceptos = []
    for l in liquidaciones:
        for k in l["calculo"]["conceptos"]:
            conceptos.append(base(l) + [l["num_id"], l["nombre"], k["nombre"],
                                        "Devengo" if k["tipo"] == "devengo" else "Deducción",
                                        k.get("cantidad"), k["valor"]])
    _hoja(wb, "Conceptos", cab_base + ["Documento", "Empleado", "Concepto", "Tipo", "Cantidad", "Valor"],
          conceptos, dinero=(10,))

    aportes = []
    for l in liquidaciones:
        a = l["calculo"]["aportes"]
        aportes.append(base(l) + [l["num_id"], l["nombre"], l["ibc"], a["salud"], a["pension"], a["arl"], a["caja"],
                                  a["sena"], a["icbf"], a["total"]])
    _hoja(wb, "Aportes empleador", cab_base + ["Documento", "Empleado", "IBC", "Salud", "Pensión", "ARL", "Caja de compensación",
                                                "SENA", "ICBF", "Total aportes"], aportes, dinero=(6, 7, 8, 9, 10, 11, 12, 13))

    prov = []
    for l in liquidaciones:
        p = l["calculo"]["provisiones"]
        prov.append(base(l) + [l["num_id"], l["nombre"], p["prima"], p["cesantias"], p["intereses"],
                               p["vacaciones"], p["total"]])
    _hoja(wb, "Provisiones", cab_base + ["Documento", "Empleado", "Prima", "Cesantías", "Intereses cesantías",
                                          "Vacaciones", "Total provisiones"], prov, dinero=(6, 7, 8, 9, 10, 11))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
