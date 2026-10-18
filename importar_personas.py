"""Carga masiva de personas desde un archivo plano (Excel .xlsx o CSV)."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import personas_db as P
from consultas import _ejecutar

MAX_FILAS = 3000
RE_CORREO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
COLUMNAS = [("nombre", "Nombre"), ("tipo_id", "Tipo de documento"), ("num_id", "Número de documento"),
            ("correo", "Correo"), ("telefono", "Teléfono"), ("cargo", "Cargo"), ("ciudad", "Ciudad"),
            ("fecha_ingreso", "Fecha de ingreso")]
# nombres de columna aceptados (sin tildes, minúsculas, sin espacios/guiones)
ALIAS = {
    "nombre": "nombre", "nombres": "nombre", "nombrecompleto": "nombre", "empleado": "nombre",
    "tipodedocumento": "tipo_id", "tipodocumento": "tipo_id", "tipoid": "tipo_id", "tipodeidentificacion": "tipo_id",
    "numerodedocumento": "num_id", "numerodocumento": "num_id", "documento": "num_id", "cedula": "num_id",
    "identificacion": "num_id", "numid": "num_id", "nodocumento": "num_id", "nit": "num_id",
    "correo": "correo", "correoelectronico": "correo", "email": "correo", "mail": "correo",
    "telefono": "telefono", "celular": "telefono", "tel": "telefono",
    "cargo": "cargo", "puesto": "cargo",
    "ciudad": "ciudad",
    "fechadeingreso": "fecha_ingreso", "fechaingreso": "fecha_ingreso", "ingreso": "fecha_ingreso",
    "fechadeingresoalaempresa": "fecha_ingreso", "fechadevinculacion": "fecha_ingreso",
}
TIPOS = {"cc": "C.C.", "cedula": "C.C.", "cedulaciudadania": "C.C.", "ceduladeciudadania": "C.C.",
         "ce": "C.E.", "cedulaextranjeria": "C.E.", "cedulaextranjera": "C.E.", "cedulade extranjeria": "C.E.",
         "ti": "T.I.", "tarjetaidentidad": "T.I.", "tarjetadeidentidad": "T.I.",
         "pep": "PEP", "ppt": "PPT", "pasaporte": "Pasaporte", "nit": "NIT"}


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _matriz(nombre, datos):
    """Lee .xlsx (todas las hojas) o CSV. Devuelve ({hoja: filas}, error)."""
    ext = (nombre or "").lower().rsplit(".", 1)[-1]
    if ext == "xlsx":
        try:
            wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}, None
        except Exception:
            return None, "No se pudo leer el Excel. Usa un archivo .xlsx válido (o guarda como CSV)."
    if ext in ("csv", "txt"):
        try:
            texto = datos.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = datos.decode("latin-1")
        muestra = texto[:2000]
        delim = ";" if muestra.count(";") > muestra.count(",") else ("\t" if muestra.count("\t") > muestra.count(",") else ",")
        return {"csv": list(csv.reader(io.StringIO(texto), delimiter=delim))}, None
    return None, "Formato no admitido. Sube un archivo .xlsx o .csv."


def filas_de(matriz, alias, requerida, columnas):
    """Convierte una matriz (1.ª fila = títulos) en filas dict. Devuelve (filas, error)."""
    matriz = [r for r in matriz if any(_txt(c) for c in r)]
    if not matriz:
        return None, "El archivo está vacío."
    cols = {}
    for i, h in enumerate(matriz[0]):
        k = alias.get(_norm(h))
        if k and k not in cols:
            cols[k] = i
    if requerida not in cols:
        return None, "No encontré la columna «%s» en la primera fila. Descarga la plantilla y úsala como guía." % requerida
    if len(matriz) - 1 > MAX_FILAS:
        return None, "El archivo tiene más de %d filas. Divídelo en partes." % MAX_FILAS
    filas = []
    for n, r in enumerate(matriz[1:], start=2):
        d = {c: _txt(r[i]) if i < len(r) else "" for c, i in cols.items()}
        for c in columnas:
            d.setdefault(c, "")
        d["_fila"] = n
        filas.append(d)
    return filas, None


def leer_archivo(nombre, datos):
    """Personas: devuelve (filas, error)."""
    hojas, error = _matriz(nombre, datos)
    if error:
        return None, error
    return filas_de(next(iter(hojas.values())), ALIAS, "nombre", [c for c, _ in COLUMNAS])


def _fecha(t):
    t = t.strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        a, mes, dia = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$", t)
        if not m:
            return None
        dia, mes, a = map(int, m.groups())
    try:
        return date(a, mes, dia)
    except ValueError:
        return None


def _limpiar(d):
    """Normaliza y valida una fila. Devuelve (datos, errores)."""
    e = []
    o = {"nombre": re.sub(r"\s+", " ", d["nombre"])[:200]}
    if len(o["nombre"]) < 3:
        e.append("Falta el nombre")
    num = re.sub(r"[\s.\-]", "", d["num_id"])[:30] if d["tipo_id"].upper() != "PASAPORTE" else re.sub(r"\s", "", d["num_id"])[:30]
    o["num_id"] = num or None
    t = d["tipo_id"]
    if t:
        o["tipo_id"] = TIPOS.get(_norm(t))
        if not o["tipo_id"]:
            e.append("Tipo de documento no reconocido «%s» (usa C.C., C.E., T.I., PEP, PPT, Pasaporte o NIT)" % t)
    else:
        o["tipo_id"] = "C.C." if num else None
    if not num:
        o["tipo_id"] = None
    o["correo"] = d["correo"].strip().lower()[:190] or None
    if o["correo"] and not RE_CORREO.match(o["correo"]):
        e.append("Correo no válido «%s»" % d["correo"])
    o["telefono"] = d["telefono"][:40] or None
    o["cargo"] = d["cargo"][:150] or None
    o["ciudad"] = d["ciudad"][:100] or None
    o["fecha_ingreso"] = None
    if d["fecha_ingreso"]:
        f = _fecha(d["fecha_ingreso"])
        if not f:
            e.append("Fecha de ingreso no válida «%s» (usa dd/mm/aaaa o aaaa-mm-dd)" % d["fecha_ingreso"])
        elif f > date.today():
            e.append("La fecha de ingreso no puede ser futura")
        else:
            o["fecha_ingreso"] = f
    return o, e


def procesar(filas, guardar=True):
    """Crea personas nuevas (sin acceso) y actualiza las existentes (por documento o correo).
    Solo se sobrescriben los campos que vienen con valor en el archivo."""
    res = {"creadas": 0, "actualizadas": 0, "sin_cambios": 0, "errores": []}
    vistos_doc, vistos_mail = {}, {}
    for d in filas:
        o, e = _limpiar(d)
        n = d["_fila"]
        if not any([d["nombre"], d["num_id"], d["correo"]]):
            continue
        if o["num_id"] in vistos_doc and o["num_id"]:
            e.append("Documento repetido en el archivo (fila %d)" % vistos_doc[o["num_id"]])
        if o["correo"] in vistos_mail and o["correo"]:
            e.append("Correo repetido en el archivo (fila %d)" % vistos_mail[o["correo"]])
        if e:
            res["errores"].append((n, d["nombre"] or "—", "; ".join(e)))
            continue
        existente = None
        if o["num_id"]:
            existente = _ejecutar("SELECT * FROM usuarios WHERE num_id=%s", [o["num_id"]], uno=True)
        if not existente and o["correo"]:
            existente = _ejecutar("SELECT * FROM usuarios WHERE correo=%s", [o["correo"]], uno=True)
        if existente:
            cambios = {c: v for c, v in o.items() if v and v != (P.a_fecha(existente[c]) if c == "fecha_ingreso" else existente[c])}
            # el correo/documento no pueden chocar con otra persona
            choque = None
            if "correo" in cambios and P.duplicado("correo", cambios["correo"], existente["id"]):
                choque = "El correo ya pertenece a otra persona"
            if "num_id" in cambios and P.duplicado("num_id", cambios["num_id"], existente["id"]):
                choque = "El documento ya pertenece a otra persona"
            if choque:
                res["errores"].append((n, o["nombre"], choque))
                continue
            if o["num_id"]: vistos_doc[o["num_id"]] = n
            if o["correo"]: vistos_mail[o["correo"]] = n
            if not cambios:
                res["sin_cambios"] += 1
            else:
                res["actualizadas"] += 1
                if guardar:
                    _ejecutar("UPDATE usuarios SET " + ", ".join(c + "=%s" for c in cambios) + " WHERE id=%s",
                              list(cambios.values()) + [existente["id"]], commit=True)
        else:
            if o["correo"] and P.duplicado("correo", o["correo"]):
                res["errores"].append((n, o["nombre"], "El correo ya existe"))
                continue
            if o["num_id"]: vistos_doc[o["num_id"]] = n
            if o["correo"]: vistos_mail[o["correo"]] = n
            res["creadas"] += 1
            if guardar:
                P.crear(o, activo=False)
    return res


def plantilla():
    wb = Workbook()
    ws = wb.active
    ws.title = "Personas"
    ws.append([t for _, t in COLUMNAS])
    ws.append(["Juan Pérez Gómez", "C.C.", "1012345678", "juan.perez@empresa.com", "3001234567", "Conductor", "Bogotá", "2021-03-15"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="254AA5")
    for col, w in zip("ABCDEFGH", (30, 20, 22, 32, 16, 24, 18, 18)):
        ws.column_dimensions[col].width = w
    for r in range(2, 3000):          # documento y teléfono como texto (conserva ceros)
        ws.cell(r, 3).number_format = "@"
        ws.cell(r, 5).number_format = "@"
        ws.cell(r, 8).number_format = "yyyy-mm-dd"
    ay = wb.create_sheet("Instrucciones")
    for l in ["Una fila por persona. La fila 1 son los títulos: no los cambies de orden ni los borres.",
              "Obligatorio: Nombre. Recomendado: Número de documento (con él se evitan duplicados).",
              "Tipo de documento: C.C., C.E., T.I., PEP, PPT, Pasaporte o NIT. Si lo dejas vacío y hay número, se asume C.C.",
              "Fecha de ingreso: la de entrada a la empresa (aaaa-mm-dd o dd/mm/aaaa). Con ella la aplicación calcula la antigüedad.",
              "Si la persona ya existe (mismo documento o mismo correo) se actualizan solo los campos que traen valor.",
              "Las personas nuevas se crean SIN acceso a la aplicación; el acceso se habilita después en Usuarios.",
              "Puedes subir este Excel o guardarlo como CSV (separado por comas o punto y coma)."]:
        ay.append([l])
    ay.column_dimensions["A"].width = 110
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
