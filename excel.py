"""Genera el archivo Excel (.xlsx) con los registros."""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import etiquetas as E

PREFIJO = [("reg_id", "ID registro"), ("reg_nombre", "Nombre / Razón social"),
           ("reg_ident", "Identificación / NIT")]
HOJAS = [
    ("Registros", E.REGISTRO_COLS, []),
    ("Representante legal", E.REPRESENTANTE_COLS, PREFIJO),
    ("Financiero", E.FINANCIERO_COLS, PREFIJO),
    ("Cuentas bancarias", E.CUENTAS_COLS, PREFIJO),
    ("Beneficiarios finales", E.BENEFICIARIOS_COLS, PREFIJO),
    ("Referencias comerciales", E.REFERENCIAS_COLS, PREFIJO),
]


def generar_excel(hojas):
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, cols, prefijo in HOJAS:
        ws = wb.create_sheet(nombre)
        todas = prefijo + cols
        ws.append([etq for _, etq in todas])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="254AA5")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        anchos = [len(etq) + 2 for _, etq in todas]
        for fila in hojas.get(nombre, []):
            vals = []
            for col, _ in todas:
                v = E.valor_plano(col, fila.get(col))
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                vals.append(v)
            ws.append(vals)
            n = ws.max_row
            for i, (col, _) in enumerate(todas, start=1):
                celda = ws.cell(row=n, column=i)
                v = celda.value
                if isinstance(v, str):
                    if v.startswith(("=", "+", "-", "@")):
                        celda.data_type = "s"      # nunca interpretar texto como fórmula
                    anchos[i - 1] = max(anchos[i - 1], min(len(v) + 2, 45))
                elif isinstance(v, datetime):
                    celda.number_format = "dd/mm/yyyy hh:mm"
                    anchos[i - 1] = max(anchos[i - 1], 18)
                elif isinstance(v, date):
                    celda.number_format = "dd/mm/yyyy"
                    anchos[i - 1] = max(anchos[i - 1], 14)
                elif col in E.COLS_DINERO:
                    celda.number_format = "#,##0"
                elif col == "participacion":
                    celda.number_format = '0.00"%"'
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(a, 10), 45)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
