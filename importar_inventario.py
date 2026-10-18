"""Carga masiva del inventario de celulares y de equipos de cómputo (con su software)."""
import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import celulares_db as C
import equipos_db as E
import personas_db as P
from consultas import _ejecutar
from importar_personas import _fecha, _matriz, _norm, _txt, filas_de


# ---------------------------------------------------------------- utilidades
def _mapa(opciones, extra=None):
    """{texto normalizado: clave} a partir de {clave: etiqueta}."""
    m = {}
    for k, v in opciones.items():
        m[_norm(k)] = k
        m[_norm(v)] = k
    for k, v in (extra or {}).items():
        m[_norm(k)] = v
    return m


def _alias(*grupos):
    a = {}
    for clave, nombres in grupos:
        for n in nombres:
            a[_norm(n)] = clave
        a[_norm(clave)] = clave
    return a


def _gb(v):
    v = (v or "").strip().lower().replace(",", ".")
    if not v:
        return None, False
    m = re.match(r"^(\d+(?:\.\d+)?)\s*(gb|tb|g|t)?$", v)
    if not m:
        return None, True
    return int(round(float(m.group(1)) * (1000 if (m.group(2) or "").startswith("t") else 1))), False


def _pesos(v):
    """'65000', '65.000', '$ 65,000', '65000,50' -> entero en pesos. None si vacío o inválido."""
    t = re.sub(r"[\s$]", "", v or "")
    if not t:
        return None
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", t):
        return int(re.sub(r"[.,]", "", t))
    t = re.sub(r"[.,]\d{1,2}$", "", t)
    return int(t) if t.isdigit() else None


class Personas:
    """Busca personas por documento, «Nombre · documento», correo o nombre exacto (si es único)."""

    def __init__(self):
        filas = _ejecutar("SELECT id, nombre, num_id, correo FROM usuarios")
        self.doc = {re.sub(r"[\s.\-]", "", f["num_id"]): f["id"] for f in filas if f["num_id"]}
        self.mail = {f["correo"].lower(): f["id"] for f in filas if f["correo"]}
        self.etq = {P.etiqueta(f).lower(): f["id"] for f in filas}
        nombres = {}
        for f in filas:
            nombres.setdefault(f["nombre"].strip().lower(), []).append(f["id"])
        self.nombre = {k: v[0] for k, v in nombres.items() if len(v) == 1}
        self.repetidos = {k for k, v in nombres.items() if len(v) > 1}

    def buscar(self, texto):
        """(id, error). texto vacío -> (None, None)."""
        t = (texto or "").strip()
        if not t:
            return None, None
        k = t.lower()
        for d in (self.etq, self.mail):
            if k in d:
                return d[k], None
        limpio = re.sub(r"[\s.\-]", "", t)
        if limpio in self.doc:
            return self.doc[limpio], None
        if k in self.nombre:
            return self.nombre[k], None
        if k in self.repetidos:
            return None, "Hay varias personas llamadas «%s»; usa el número de documento" % t
        return None, "La persona «%s» no existe en Usuarios (cárgala primero)" % t


def _resultado():
    return {"creadas": 0, "actualizadas": 0, "sin_cambios": 0, "errores": [], "extra": []}


# ---------------------------------------------------------------- CELULARES
COL_CEL = [  # clave, título, ejemplo
    ("marca", "Marca", "Samsung"), ("modelo", "Modelo", "Galaxy A54"),
    ("imei1", "IMEI 1", "356938035643809"), ("imei2", "IMEI 2", ""), ("sn", "SN", "R5CT12ABCDE"),
    ("fecha_compra", "Fecha de compra", "2024-02-15"), ("acta", "Firma / Acta", "Firmada"),
    ("linea", "Línea asignada", "3001234567"), ("operador", "Operador", "Claro"),
    ("valor_plan", "Valor plan", "65000"),
    ("usuario_actual", "Usuario actual", "1012345678"), ("usuario_anterior", "Usuario anterior", ""),
    ("estado", "Estado", "Activo"), ("baja_fecha", "Fecha de baja", ""), ("baja_motivo", "Motivo de baja", ""),
    ("observaciones", "Observaciones", ""),
]
ALIAS_CEL = _alias(
    ("marca", ["marca"]), ("modelo", ["modelo"]), ("imei1", ["imei1", "imei", "imei 1"]), ("imei2", ["imei2", "imei 2"]),
    ("sn", ["sn", "serial", "serie", "numerodeserie"]),
    ("fecha_compra", ["fechacompra", "fechadecompra", "compra"]),
    ("acta", ["firmaacta", "firma/acta", "acta", "firma"]),
    ("linea", ["lineaasignada", "linea", "numero", "linea asignada"]), ("operador", ["operador"]),
    ("valor_plan", ["valorplan", "valordelplan", "plan"]),
    ("usuario_actual", ["usuarioactual", "usuario", "asignadoa", "actual"]),
    ("usuario_anterior", ["usuarioanterior", "anterior"]),
    ("estado", ["estado", "situacion"]), ("baja_fecha", ["fechabaja", "fechadebaja"]),
    ("baja_motivo", ["motivobaja", "motivodebaja"]), ("observaciones", ["observaciones", "observacion", "notas"]),
)
ACTAS = _mapa(C.ACTAS, {"si": "firmada", "no": "pendiente"})
MOTIVOS = _mapa(C.MOTIVOS_BAJA, {"dañado": "dano", "danado": "dano", "robado": "robo", "perdido": "perdida"})


def leer_celulares(nombre, datos):
    hojas, error = _matriz(nombre, datos)
    if error:
        return None, error
    return filas_de(next(iter(hojas.values())), ALIAS_CEL, "marca", [c for c, _, _ in COL_CEL])


def _limpiar_cel(d, pers):
    e = []
    o = {"marca": d["marca"].strip()[:60], "modelo": d["modelo"].strip()[:80]}
    if not o["marca"]:
        e.append("Falta la marca")
    if not o["modelo"]:
        e.append("Falta el modelo")
    for c, n in (("imei1", "IMEI 1"), ("imei2", "IMEI 2")):
        v = re.sub(r"\D", "", d[c])
        o[c] = v or None
        if v and not re.fullmatch(r"\d{14,16}", v):
            e.append("%s debe tener entre 14 y 16 dígitos" % n)
    o["sn"] = re.sub(r"\s", "", d["sn"])[:40] or None
    if not o["imei1"] and not o["sn"]:
        e.append("Cada celular necesita IMEI 1 o SN (así se evita duplicar)")
    if o["imei1"] and o["imei1"] == o["imei2"]:
        e.append("IMEI 1 e IMEI 2 no pueden ser iguales")
    o["fecha_compra"] = None
    if d["fecha_compra"]:
        f = _fecha(d["fecha_compra"])
        if not f:
            e.append("Fecha de compra no válida «%s»" % d["fecha_compra"])
        elif f > date.today():
            e.append("La fecha de compra no puede ser futura")
        else:
            o["fecha_compra"] = f
    o["acta"] = ACTAS.get(_norm(d["acta"])) if d["acta"] else None
    if d["acta"] and not o["acta"]:
        e.append("Firma/Acta debe ser Firmada, Pendiente o No aplica")
    lin = re.sub(r"[^\d+]", "", d["linea"])
    o["linea"] = lin or None
    if lin and not re.fullmatch(r"\+?\d{7,20}", lin):
        e.append("La línea asignada solo debe tener números")
    o["operador"] = d["operador"].strip()[:40] or None
    o["valor_plan"] = _pesos(d["valor_plan"])
    if d["valor_plan"] and o["valor_plan"] is None:
        e.append("Valor plan no válido «%s»" % d["valor_plan"])
    o["observaciones"] = d["observaciones"].strip()[:2000] or None
    o["persona_actual_id"], err = pers.buscar(d["usuario_actual"])
    if err:
        e.append("Usuario actual: " + err)
    o["persona_anterior_id"], err = pers.buscar(d["usuario_anterior"])
    if err:
        e.append("Usuario anterior: " + err)
    est = _norm(d["estado"])
    o["baja"] = est in ("baja", "dadodebaja", "dadabaja", "inactivo")
    if est and not o["baja"] and est not in ("activo", "disponible", "asignado"):
        e.append("Estado debe ser Activo o Baja")
    o["baja_fecha"], o["baja_motivo"] = None, None
    if o["baja"]:
        f = _fecha(d["baja_fecha"]) if d["baja_fecha"] else date.today()
        if not f or f > date.today():
            e.append("Fecha de baja no válida")
        o["baja_fecha"] = f
        o["baja_motivo"] = MOTIVOS.get(_norm(d["baja_motivo"])) or "otro"
        if d["baja_motivo"] and not MOTIVOS.get(_norm(d["baja_motivo"])):
            e.append("Motivo de baja debe ser Daño, Pérdida, Robo, Obsolescencia, Venta u Otro")
        if o["persona_actual_id"]:
            o["persona_anterior_id"] = o["persona_anterior_id"] or o["persona_actual_id"]
            o["persona_actual_id"] = None
    return o, e


def procesar_celulares(filas, quien, guardar=True):
    res, pers = _resultado(), Personas()
    visto = set()
    hoy = date.today().isoformat()
    for d in filas:
        if not any(d[c] for c in ("marca", "modelo", "imei1", "sn")):
            continue
        n, o, e = d["_fila"], *_limpiar_cel(d, pers)
        for k in ("imei1", "sn"):
            if o[k] and (k, o[k]) in visto:
                e.append("%s repetido en el archivo" % ("IMEI 1" if k == "imei1" else "SN"))
        if e:
            res["errores"].append(("Celulares", n, "%s %s" % (d["marca"], d["modelo"]), "; ".join(e)))
            continue
        ex = None
        if o["imei1"]:
            ex = _ejecutar("SELECT * FROM celulares WHERE imei1=%s", [o["imei1"]], uno=True)
        if not ex and o["sn"]:
            ex = _ejecutar("SELECT * FROM celulares WHERE sn=%s", [o["sn"]], uno=True)
        for k in ("imei1", "sn"):
            if o[k]:
                visto.add((k, o[k]))
        etiqueta = "%s %s" % (o["marca"], o["modelo"])
        if ex:
            cambios = {}
            for c in C.CAMPOS:
                v = o.get(c)
                act = C.a_fecha(ex[c]) if c == "fecha_compra" else (int(float(ex[c] or 0)) if c == "valor_plan" else ex[c])
                if v not in (None, "") and v != act:
                    cambios[c] = v
            choque = None
            for k in ("imei1", "sn"):
                if k in cambios and C.duplicado(k, cambios[k], ex["id"]):
                    choque = "%s ya pertenece a otro celular" % ("IMEI 1" if k == "imei1" else "SN")
            if choque:
                res["errores"].append(("Celulares", n, etiqueta, choque))
                continue
            cid = ex["id"]
            mover = (not o["baja"] and ex["estado"] == "activo" and o["persona_actual_id"]
                     and o["persona_actual_id"] != ex["persona_actual_id"])
            dar_baja = o["baja"] and ex["estado"] != "baja"
            if not (cambios or mover or dar_baja):
                res["sin_cambios"] += 1
                continue
            res["actualizadas"] += 1
            if guardar:
                if cambios:
                    _ejecutar("UPDATE celulares SET " + ", ".join(c + "=%s" for c in cambios) + " WHERE id=%s",
                              list(cambios.values()) + [cid], commit=True)
                if mover:
                    C.reasignar(cid, o["persona_actual_id"], hoy, "Carga masiva", quien)
                if dar_baja:
                    C.dar_baja(cid, o["baja_fecha"].isoformat(), o["baja_motivo"], "Carga masiva", quien)
        else:
            res["creadas"] += 1
            if guardar:
                datos = {c: o.get(c) for c in C.CAMPOS}
                datos["acta"] = o["acta"] or "pendiente"
                datos["valor_plan"] = o["valor_plan"] or 0
                cid = C.crear(datos, o["persona_actual_id"], hoy, quien)
                if o["persona_anterior_id"]:
                    _ejecutar("UPDATE celulares SET persona_anterior_id=%s WHERE id=%s",
                              [o["persona_anterior_id"], cid], commit=True)
                if o["baja"]:
                    C.dar_baja(cid, o["baja_fecha"].isoformat(), o["baja_motivo"], "Carga masiva", quien)
    return res


# ---------------------------------------------------------------- EQUIPOS
COL_EQ = [
    ("codigo", "ID de inventario", "EQ-0001"), ("tipo", "Tipo de equipo", "Portátil"), ("marca", "Marca", "Dell"),
    ("modelo", "Modelo", "Latitude 5440"), ("serie", "Serie", "ABC1234"),
    ("procesador", "Procesador", "Intel Core i5-1335U"), ("ram_gb", "Memoria RAM (GB)", "16"),
    ("ssd_gb", "Almacenamiento SSD (GB)", "512"), ("usuario", "Usuario asignado", "1012345678"),
    ("ciudad", "Ciudad", "Bogotá"), ("sistema_operativo", "Sistema operativo", "Windows 11 Pro"),
    ("office", "Microsoft Office", "Microsoft 365"), ("antivirus", "Antivirus", "ESET"),
    ("estado", "Estado", "Activo"), ("observaciones", "Observaciones", ""),
]
COL_SW = [("codigo", "ID de inventario", "EQ-0001"), ("nombre", "Programa", "AutoCAD LT"), ("version", "Versión", "2025"),
          ("tipo_licencia", "Tipo de licencia", "Suscripción"), ("estado", "Estado", "Activo"),
          ("vencimiento", "Vencimiento", "2027-01-06"), ("notas", "Notas", "")]
ALIAS_EQ = _alias(
    ("codigo", ["idinventario", "iddeinventario", "id", "codigo", "idequipo"]),
    ("tipo", ["tipodeequipo", "tipoequipo", "tipo"]), ("marca", ["marca"]), ("modelo", ["modelo"]),
    ("serie", ["serie", "serial", "sn", "numerodeserie"]), ("procesador", ["procesador", "cpu"]),
    ("ram_gb", ["memoriaram", "ram", "memoriaramgb", "ramgb"]),
    ("ssd_gb", ["almacenamientossd", "ssd", "almacenamiento", "almacenamientossdgb", "ssdgb", "discoduro"]),
    ("usuario", ["usuarioasignado", "usuario", "asignadoa"]), ("ciudad", ["ciudad"]),
    ("sistema_operativo", ["sistemaoperativo", "so"]), ("office", ["microsoftoffice", "office"]),
    ("antivirus", ["antivirus"]), ("estado", ["estado"]), ("observaciones", ["observaciones", "observacion", "notas"]),
)
ALIAS_SW = _alias(
    ("codigo", ["idinventario", "iddeinventario", "id", "codigo", "idequipo"]),
    ("nombre", ["programa", "software", "nombre", "aplicacion"]), ("version", ["version"]),
    ("tipo_licencia", ["tipodelicencia", "tipolicencia", "licencia"]), ("estado", ["estado"]),
    ("vencimiento", ["vencimiento", "fechadevencimiento", "vence"]), ("notas", ["notas", "observaciones"]),
)
TIPOS_EQ = _mapa(E.TIPOS, {"laptop": "portatil", "portatiles": "portatil", "computador": "escritorio", "pc": "escritorio",
                           "desktop": "escritorio", "escritoriocomputador": "escritorio", "allinone": "todo_en_uno",
                           "aio": "todo_en_uno", "tecladoymouse": "teclado_mouse", "combo": "teclado_mouse",
                           "camaraweb": "camara", "webcam": "camara", "pantalla": "monitor"})
ESTADOS_EQ = _mapa(E.ESTADOS, {"dadodebaja": "baja", "enreparacion": "reparacion", "bodega": "almacenado"})
LICENCIAS_SW = _mapa(E.LICENCIAS)
ESTADOS_SW = _mapa(E.SW_ESTADOS)


def leer_equipos(nombre, datos):
    """Devuelve ({'equipos': filas, 'software': filas}, error). Un CSV con columna «Programa» es solo software."""
    hojas, error = _matriz(nombre, datos)
    if error:
        return None, error
    eq, sw = None, None
    for titulo, m in hojas.items():
        cabecera = {ALIAS_SW.get(_norm(h)) for h in (m[0] if m else [])}
        if "nombre" in cabecera and "codigo" in cabecera and ("tipo" not in {ALIAS_EQ.get(_norm(h)) for h in m[0]}):
            sw, err = filas_de(m, ALIAS_SW, "nombre", [c for c, _, _ in COL_SW])
            if err:
                return None, "Hoja «%s»: %s" % (titulo, err)
        elif eq is None:
            eq, err = filas_de(m, ALIAS_EQ, "tipo", [c for c, _, _ in COL_EQ])
            if err:
                return None, "Hoja «%s»: %s" % (titulo, err)
    if eq is None and sw is None:
        return None, "No encontré columnas de equipos (Tipo de equipo) ni de software (Programa). Usa la plantilla."
    return {"equipos": eq or [], "software": sw or []}, None


def _limpiar_eq(d, pers):
    e = []
    o = {"codigo": re.sub(r"\s", "", d["codigo"])[:30] or None}
    o["tipo"] = TIPOS_EQ.get(_norm(d["tipo"])) if d["tipo"] else None
    if d["tipo"] and not o["tipo"]:
        e.append("Tipo de equipo no reconocido «%s» (%s)" % (d["tipo"], ", ".join(E.TIPOS.values())))
    for c, n in (("marca", 60), ("modelo", 80), ("serie", 60), ("procesador", 100), ("ciudad", 100),
                 ("sistema_operativo", 100), ("office", 100), ("antivirus", 100)):
        o[c] = d[c].strip()[:n] or None
    o["observaciones"] = d["observaciones"].strip()[:2000] or None
    for c, n in (("ram_gb", "RAM"), ("ssd_gb", "SSD")):
        o[c], mal = _gb(d[c])
        if mal:
            e.append("%s debe estar en GB (16) o TB (1 TB)" % n)
    o["estado"] = ESTADOS_EQ.get(_norm(d["estado"])) if d["estado"] else None
    if d["estado"] and not o["estado"]:
        e.append("Estado debe ser Activo, En reparación, Almacenado o Dado de baja")
    o["persona_id"], err = pers.buscar(d["usuario"])
    if err:
        e.append("Usuario asignado: " + err)
    if o["tipo"] and o["tipo"] not in E.COMPUTO:
        for c in ("procesador", "ram_gb", "ssd_gb", "sistema_operativo", "office", "antivirus"):
            o[c] = None
    return o, e


def procesar_equipos(paquete, guardar=True):
    res, pers = _resultado(), Personas()
    res["extra"] = []
    vistos = {}
    codigos = {}                       # codigo -> id (existentes o creados en esta carga)
    for d in paquete["equipos"]:
        if not any(d[c] for c in ("codigo", "tipo", "marca", "modelo", "serie")):
            continue
        n, (o, e) = d["_fila"], _limpiar_eq(d, pers)
        if o["codigo"] and o["codigo"].lower() in vistos:
            e.append("ID de inventario repetido en el archivo (fila %d)" % vistos[o["codigo"].lower()])
        ex = None
        if o["codigo"]:
            ex = _ejecutar("SELECT * FROM equipos WHERE codigo=%s", [o["codigo"]], uno=True)
        if not ex and o["serie"] and not o["codigo"]:
            ex = _ejecutar("SELECT * FROM equipos WHERE serie=%s", [o["serie"]], uno=True)
        if not ex and not o["tipo"]:
            e.append("Falta el tipo de equipo")
        if e:
            res["errores"].append(("Equipos", n, o["codigo"] or d["marca"] or "—", "; ".join(e)))
            continue
        if o["codigo"]:
            vistos[o["codigo"].lower()] = n
        if ex:
            cambios = {c: o[c] for c in E.CAMPOS if c != "codigo" and o.get(c) not in (None, "") and o[c] != ex[c]}
            if not cambios:
                res["sin_cambios"] += 1
            else:
                res["actualizadas"] += 1
                if guardar:
                    _ejecutar("UPDATE equipos SET " + ", ".join(c + "=%s" for c in cambios) + " WHERE id=%s",
                              list(cambios.values()) + [ex["id"]], commit=True)
            codigos[ex["codigo"].lower()] = ex["id"]
        else:
            res["creadas"] += 1
            nuevo = dict(o)
            nuevo["estado"] = o["estado"] or "activo"
            if not nuevo["codigo"]:
                nuevo["codigo"] = E.siguiente_codigo() if guardar else "(automático)"
            if guardar:
                codigos[nuevo["codigo"].lower()] = E.crear(nuevo)
            else:
                codigos[nuevo["codigo"].lower()] = -1
    # ---- software
    sw_nuevos = sw_act = 0
    for d in paquete["software"]:
        if not (d["nombre"] or d["codigo"]):
            continue
        n, e = d["_fila"], []
        cod = re.sub(r"\s", "", d["codigo"]).lower()
        eid = codigos.get(cod)
        if eid is None and cod:
            r = _ejecutar("SELECT id FROM equipos WHERE codigo=%s", [d["codigo"].strip()], uno=True)
            eid = r["id"] if r else None
        if not cod:
            e.append("Falta el ID de inventario del equipo")
        elif eid is None:
            e.append("El equipo «%s» no existe (créalo en la hoja de equipos)" % d["codigo"])
        if not d["nombre"].strip():
            e.append("Falta el nombre del programa")
        lic = ESTADOS_SW.get(_norm(d["estado"])) if d["estado"] else "activo"
        if not lic:
            e.append("Estado del software debe ser Activo, Por renovar, Vencido, Sin licencia o Desinstalado")
        tl = LICENCIAS_SW.get(_norm(d["tipo_licencia"])) if d["tipo_licencia"] else None
        if d["tipo_licencia"] and not tl:
            e.append("Tipo de licencia no reconocido «%s»" % d["tipo_licencia"])
        venc = None
        if d["vencimiento"]:
            venc = _fecha(d["vencimiento"])
            if not venc:
                e.append("Vencimiento no válido «%s»" % d["vencimiento"])
        if e:
            res["errores"].append(("Software", n, d["nombre"] or "—", "; ".join(e)))
            continue
        nombre, version = d["nombre"].strip()[:150], d["version"].strip()[:50] or None
        ya = None
        if eid and eid > 0:
            ya = _ejecutar("SELECT id FROM equipo_software WHERE equipo_id=%s AND LOWER(nombre)=%s AND "
                           "COALESCE(version,'')=%s", [eid, nombre.lower(), version or ""], uno=True)
        v = [tl, lic, venc, d["notas"].strip()[:300] or None]
        if ya:
            sw_act += 1
            if guardar:
                _ejecutar("UPDATE equipo_software SET tipo_licencia=%s, estado=%s, vencimiento=%s, notas=%s WHERE id=%s",
                          v + [ya["id"]], commit=True)
        else:
            sw_nuevos += 1
            if guardar:
                _ejecutar("INSERT INTO equipo_software (equipo_id, nombre, version, tipo_licencia, estado, vencimiento, "
                          "notas) VALUES (%s, %s, %s, %s, %s, %s, %s)", [eid, nombre, version] + v, commit=True)
    if paquete["software"]:
        res["extra"].append("Software: %d nuevos, %d actualizados" % (sw_nuevos, sw_act))
    return res


# ---------------------------------------------------------------- PLANTILLAS
def _hoja(wb, titulo, cols, ejemplos, anchos=None, primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo
    ws.append([t for _, t, _ in cols])
    for fila in ejemplos:
        ws.append(fila)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="254AA5")
    for i in range(len(cols)):
        ws.column_dimensions[chr(65 + i)].width = max(14, len(cols[i][1]) + 4)
    return ws


def _lista(ws, col, opciones, filas=3000):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(opciones), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col, col, filas))


def _texto(ws, cols, claves, filas=3000):
    for r in range(2, filas):
        for i, (k, _, _) in enumerate(cols):
            if k in claves:
                ws.cell(r, i + 1).number_format = "@"


def _guardar(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _instrucciones(wb, lineas):
    ay = wb.create_sheet("Instrucciones")
    for l in lineas:
        ay.append([l])
    ay.column_dimensions["A"].width = 120


def plantilla_celulares():
    wb = Workbook()
    ws = _hoja(wb, "Celulares", COL_CEL, [
        [e for _, _, e in COL_CEL],
        ["Apple", "iPhone 13", "353456789012345", "", "F2LXYZ12345", "2023-06-10", "Pendiente", "3109876543", "Movistar",
         "80000", "Ana Gómez", "1012345678", "Activo", "", "", "Rota con funda"]], primera=True)
    _texto(ws, COL_CEL, {"imei1", "imei2", "sn", "linea", "usuario_actual", "usuario_anterior"})
    idx = {k: chr(65 + i) for i, (k, _, _) in enumerate(COL_CEL)}
    _lista(ws, idx["acta"], ["Firmada", "Pendiente", "No aplica"])
    _lista(ws, idx["estado"], ["Activo", "Baja"])
    _lista(ws, idx["baja_motivo"], list(C.MOTIVOS_BAJA.values()))
    _instrucciones(wb, [
        "Una fila por celular. La fila 1 son los títulos: no los cambies. Borra las filas de ejemplo antes de cargar.",
        "Obligatorios: Marca, Modelo y IMEI 1 o SN (con eso se reconoce el celular y no se duplica).",
        "Usuario actual / Usuario anterior: escribe el número de documento (o el nombre exacto) de una persona que YA exista en Usuarios.",
        "Estado: Activo o Baja. Si es Baja, llena Fecha de baja y Motivo (Daño, Pérdida, Robo, Obsolescencia, Venta, Otro).",
        "Firma / Acta: Firmada, Pendiente o No aplica. Fechas: aaaa-mm-dd o dd/mm/aaaa. Valor plan: solo números (pesos).",
        "Si el celular ya existe (mismo IMEI 1 o SN) se actualizan los campos con valor; si cambia el Usuario actual se registra la rotación en el historial.",
        "La antigüedad y los años se calculan solos desde la Fecha de compra.",
    ])
    return _guardar(wb)


def plantilla_equipos():
    wb = Workbook()
    ws = _hoja(wb, "Equipos", COL_EQ, [
        [e for _, _, e in COL_EQ],
        ["EQ-0002", "Monitor", "LG", "24MK600", "SER998877", "", "", "", "Juan Pérez", "Medellín", "", "", "", "Activo", ""]],
        primera=True)
    _texto(ws, COL_EQ, {"codigo", "serie", "usuario"})
    idx = {k: chr(65 + i) for i, (k, _, _) in enumerate(COL_EQ)}
    _lista(ws, idx["tipo"], list(E.TIPOS.values()))
    _lista(ws, idx["estado"], list(E.ESTADOS.values()))
    sw = _hoja(wb, "Software", COL_SW, [[e for _, _, e in COL_SW]])
    _texto(sw, COL_SW, {"codigo"})
    _lista(sw, "D", list(E.LICENCIAS.values()))
    _lista(sw, "E", list(E.SW_ESTADOS.values()))
    _instrucciones(wb, [
        "Hoja «Equipos»: una fila por equipo. Hoja «Software»: una fila por programa instalado (varios por equipo).",
        "Borra las filas de ejemplo antes de cargar. No cambies los títulos de la fila 1.",
        "ID de inventario: si lo dejas vacío en un equipo nuevo, la aplicación asigna el siguiente (EQ-0001, EQ-0002…).",
        "En «Software», el ID de inventario debe ser el del equipo (de la hoja Equipos o uno ya existente).",
        "Tipos: " + ", ".join(E.TIPOS.values()) + ". Los periféricos no llevan procesador, RAM, SSD, SO, Office ni antivirus.",
        "RAM y SSD en GB (16, 512) o TB (1 TB). Usuario asignado: documento o nombre exacto de una persona que YA exista en Usuarios.",
        "Estado del equipo: " + ", ".join(E.ESTADOS.values()) + ". Estado del software: " + ", ".join(E.SW_ESTADOS.values()) + ".",
        "Tipo de licencia: " + ", ".join(E.LICENCIAS.values()) + ". Vencimiento: aaaa-mm-dd o dd/mm/aaaa.",
        "Si el equipo ya existe (mismo ID) se actualizan los campos con valor; el software se reconoce por equipo + programa + versión.",
        "Si subes CSV: un archivo para equipos y otro (con la columna «Programa») para el software.",
    ])
    return _guardar(wb)
